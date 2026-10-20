import re
from datetime import datetime, date
from pathlib import Path

import psycopg
from openpyxl import load_workbook

EXCEL = Path(r"c:\Users\jesus\Downloads\FAU-CIDWO DATA GENERAL.xlsx")
ANIO = re.compile(r"^(1[5-9]\d{2}|20\d{2})$")


def texto(valor):
    if valor is None:
        return ""
    if isinstance(valor, datetime):
        if valor.month == 1 and valor.day == 1 and 1500 <= valor.year <= 2030:
            return str(valor.year)
        return valor.strftime("%Y-%m-%d")
    if isinstance(valor, date):
        return valor.isoformat()
    if isinstance(valor, float) and valor.is_integer():
        return str(int(valor))
    return str(valor).replace("\xa0", " ").strip()


def recortar(valor, limite):
    valor = " ".join(valor.split())
    if len(valor) <= limite:
        return valor, False
    return valor[:limite], True


def fecha_de(valor):
    original = texto(valor)
    if ANIO.fullmatch(original):
        return int(original), original, False
    return None, original, True


def partir_personas(nombre, apellido):
    nombres = [p.strip() for p in re.split(r"\s*/\s*|\s*;\s*", nombre) if p.strip()] if nombre else []
    apellidos = [p.strip() for p in re.split(r"\s*/\s*|\s*;\s*", apellido) if p.strip()] if apellido else []
    if len(nombres) > 1 and len(apellidos) == len(nombres):
        return list(zip(nombres, apellidos))
    if len(nombres) > 1:
        ape = apellidos[0] if len(apellidos) == 1 else ""
        return [(n, ape) for n in nombres]
    if len(apellidos) > 1:
        nom = nombres[0] if len(nombres) == 1 else ""
        return [(nom, a) for a in apellidos]
    return [(nombre, apellido)]


def nombre_completo(nombre, apellido):
    if nombre and apellido and nombre.lower() == apellido.lower():
        return nombre
    return " ".join(p for p in (nombre, apellido) if p).strip()


def copias_de(valor):
    crudo = texto(valor)
    if not crudo:
        return None, crudo
    if re.fullmatch(r"(?i)u[nñ]ico", crudo):
        return [1], crudo
    if re.fullmatch(r"\d+", crudo):
        cantidad = int(crudo)
        if 1 <= cantidad <= 20:
            return list(range(1, cantidad + 1)), crudo
        return None, crudo
    if re.fullmatch(r"\d+(?:\.\d+)+", crudo):
        numeros = [int(p) for p in crudo.split(".")]
        if numeros and all(1 <= n <= 30 for n in numeros):
            return numeros, crudo
    partes = re.split(r"\s*[,;/]\s*|\s+y\s+", crudo)
    numeros = [int(p) for p in partes if re.fullmatch(r"\d+", p)]
    if numeros and len(numeros) == len([p for p in partes if p]):
        return numeros, crudo
    return None, crudo


def ubicacion_de(valor):
    original = " ".join(texto(valor).split())
    if not original or original in {"?", "-", "??"}:
        return original, None, None, True
    if "/" in original:
        estanteria, posicion = original.split("/", 1)
        estanteria, corto_e = recortar(estanteria, 20)
        posicion, corto_p = recortar(posicion, 20)
        return original, estanteria, posicion, corto_e or corto_p
    revisa = re.search(r"(?i)est|\d", original) is None
    estanteria, corto = recortar(original, 20)
    return original, estanteria, None, revisa or corto or len(original) > 20


def main():
    libro = load_workbook(EXCEL, data_only=True, read_only=True)
    hoja = libro.active
    filas = []
    vistas = set()
    for indice, fila in enumerate(hoja.iter_rows(values_only=True), start=1):
        if indice == 1:
            continue
        titulo = texto(fila[1] if len(fila) > 1 else None)
        if not titulo:
            continue
        clave = tuple(
            " ".join(texto(fila[columna] if len(fila) > columna else None).split()).casefold()
            for columna in (1, 2, 3, 4, 5, 6, 7, 8, 9)
        )
        if clave in vistas:
            continue
        vistas.add(clave)
        filas.append(
            {
                "titulo": titulo,
                "nombre": texto(fila[2] if len(fila) > 2 else None),
                "apellido": texto(fila[3] if len(fila) > 3 else None),
                "fecha": fila[4] if len(fila) > 4 else None,
                "tipo": texto(fila[5] if len(fila) > 5 else None),
                "fila": fila[6] if len(fila) > 6 else None,
                "editorial": texto(fila[7] if len(fila) > 7 else None),
                "cota": texto(fila[8] if len(fila) > 8 else None),
                "ejemplares": fila[9] if len(fila) > 9 else None,
            }
        )

    props = Path(__file__).resolve().parents[1].joinpath("src/main/resources/application.properties").read_text(encoding="utf-8")
    clave_bd = next(
        linea.split("=", 1)[1].strip()
        for linea in props.splitlines()
        if linea.startswith("spring.datasource.password=")
    )
    conn = psycopg.connect(
        host="127.0.0.1",
        port=5432,
        dbname="biblioteca_fau",
        user="postgres",
        password=clave_bd,
    )
    cur = conn.cursor()
    cur.execute(
        """
        TRUNCATE TABLE
            incidencias_datos,
            prestamos,
            ejemplares,
            libro_autor,
            libros,
            autores,
            editoriales,
            ubicaciones
        RESTART IDENTITY CASCADE
        """
    )

    editoriales = {}
    autores = {}
    ubicaciones = {}
    incidencias = 0
    ejemplares = 0

    def editorial_id(nombre):
        limpio, _ = recortar(nombre, 255)
        if not limpio:
            return None
        clave = limpio.lower()
        if clave not in editoriales:
            cur.execute(
                "INSERT INTO editoriales (nombre) VALUES (%s) RETURNING id_editorial",
                (limpio,),
            )
            editoriales[clave] = cur.fetchone()[0]
        return editoriales[clave]

    def autor_id(nombre, apellido):
        nombre, _ = recortar(nombre, 150)
        apellido, _ = recortar(apellido, 150)
        if not nombre and apellido:
            nombre = apellido
            apellido = ""
        if not nombre:
            nombre = "Sin autor"
        completo, _ = recortar(nombre_completo(nombre, apellido), 300)
        clave = (nombre.lower(), apellido.lower())
        if clave not in autores:
            cur.execute(
                """
                INSERT INTO autores (nombre, apellido, nombre_completo)
                VALUES (%s, %s, %s)
                RETURNING id_autor
                """,
                (nombre, apellido or None, completo or None),
            )
            autores[clave] = cur.fetchone()[0]
        return autores[clave]

    def ubicacion_id(valor):
        original, estanteria, posicion, revisa = ubicacion_de(valor)
        original, cortado = recortar(original, 100)
        if not original:
            return None, True
        revisa = revisa or cortado
        clave = original.lower()
        if clave not in ubicaciones:
            cur.execute(
                """
                INSERT INTO ubicaciones
                    (ubicacion_original, estanteria, posicion, requiere_revision)
                VALUES (%s, %s, %s, %s)
                RETURNING id_ubicacion
                """,
                (original, estanteria, posicion, revisa),
            )
            ubicaciones[clave] = cur.fetchone()[0]
        return ubicaciones[clave], revisa

    def incidencia(id_libro, campo, valor, descripcion):
        nonlocal incidencias
        cur.execute(
            """
            INSERT INTO incidencias_datos
                (id_libro, campo, valor_original, descripcion, resuelto)
            VALUES (%s, %s, %s, %s, FALSE)
            """,
            (id_libro, campo[:100], (valor or "")[:255] or None, descripcion),
        )
        incidencias += 1

    for fila in filas:
        titulo, titulo_cortado = recortar(fila["titulo"], 500)
        anio, fecha_original, revisar_fecha = fecha_de(fila["fecha"])
        fecha_original, fecha_cortada = recortar(fecha_original, 50)
        tipo, _ = recortar(fila["tipo"], 100)
        cota, cota_cortada = recortar(fila["cota"], 100)
        id_editorial = editorial_id(fila["editorial"])
        cur.execute(
            """
            INSERT INTO libros (
                titulo, id_editorial, tipo, fecha_publicacion, fecha_original,
                requiere_revision_fecha, cota
            )
            VALUES (%s, %s, %s, %s, %s, %s, %s)
            RETURNING id_libro
            """,
            (
                titulo,
                id_editorial,
                tipo or None,
                anio,
                fecha_original or None,
                revisar_fecha or fecha_cortada,
                cota or None,
            ),
        )
        id_libro = cur.fetchone()[0]

        if titulo_cortado:
            incidencia(id_libro, "TITULO", fila["titulo"], "El título superaba el largo permitido y se conservó recortado.")
        if revisar_fecha or not fecha_original:
            incidencia(id_libro, "FECHA", fecha_original, "La fecha está vacía o no es un año válido. Se conservó el valor original.")
        if not fila["editorial"]:
            incidencia(id_libro, "EDITORIAL", "", "La editorial está vacía.")
        if not cota or cota in {"?", "-", "??"} or cota_cortada:
            incidencia(id_libro, "COTA", cota, "La cota está vacía, es ambigua o fue recortada.")
        if not tipo:
            incidencia(id_libro, "TIPO", "", "El tipo de material está vacío.")

        personas = partir_personas(fila["nombre"], fila["apellido"])
        if not fila["nombre"] and not fila["apellido"]:
            incidencia(id_libro, "AUTOR", "", "El autor está vacío.")
        vistos = set()
        for nombre, apellido in personas:
            id_autor = autor_id(nombre, apellido)
            if id_autor in vistos:
                continue
            vistos.add(id_autor)
            cur.execute(
                "INSERT INTO libro_autor (id_libro, id_autor) VALUES (%s, %s) ON CONFLICT DO NOTHING",
                (id_libro, id_autor),
            )

        id_ubicacion, ubicacion_revisa = ubicacion_id(fila["fila"])
        if id_ubicacion is None or ubicacion_revisa:
            incidencia(id_libro, "FILA", texto(fila["fila"]), "La ubicación está vacía o no tiene el formato de estantería esperado.")

        numeros, crudo_copias = copias_de(fila["ejemplares"])
        if not numeros:
            incidencia(id_libro, "Nº.EJEMPLARES", crudo_copias, "La cantidad de ejemplares está vacía o no se pudo interpretar. Se registró una copia para revisión.")
            numeros = [None]
            revision_copia = True
        else:
            revision_copia = False

        for numero in numeros:
            cur.execute(
                """
                INSERT INTO ejemplares
                    (numero_ejemplar, disponible, requiere_revision, id_libro, id_ubicacion)
                VALUES (%s, TRUE, %s, %s, %s)
                """,
                (numero, revision_copia or ubicacion_revisa, id_libro, id_ubicacion),
            )
            ejemplares += 1

    conn.commit()
    cur.execute("SELECT COUNT(*) FROM libros")
    print("LIBROS", cur.fetchone()[0])
    cur.execute("SELECT COUNT(*) FROM autores")
    print("AUTORES", cur.fetchone()[0])
    cur.execute("SELECT COUNT(*) FROM editoriales")
    print("EDITORIALES", cur.fetchone()[0])
    cur.execute("SELECT COUNT(*) FROM ubicaciones")
    print("UBICACIONES", cur.fetchone()[0])
    print("EJEMPLARES", ejemplares)
    print("INCIDENCIAS", incidencias)
    conn.close()


if __name__ == "__main__":
    main()

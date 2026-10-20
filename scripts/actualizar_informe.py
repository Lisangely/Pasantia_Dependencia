from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph

ORIGEN = Path(r"c:\Users\jesus\Downloads\informe (1).docx")
DESTINO = Path(__file__).resolve().parents[1] / "docs" / "Informe_Biblioteca_FAU.docx"

REEMPLAZOS = [
    ("la bibliotecaria y los ayudantes de biblioteca", "el administrador y los ayudantes de biblioteca"),
    ("la bibliotecaria y los ayudantes", "el administrador y los ayudantes"),
    ("exclusivo de la bibliotecaria y los ayudantes", "exclusivo del administrador y los ayudantes"),
    ("por la bibliotecaria y los ayudantes", "por el administrador y los ayudantes"),
    ("correspondientes a la bibliotecaria y los ayudantes de biblioteca", "correspondientes al administrador y los ayudantes de biblioteca"),
    ("(bibliotecaria y ayudantes)", "(administrador y ayudantes)"),
    ("Bibliotecaria.", "Administrador."),
    ("bibliotecaria o ayudante", "administrador o ayudante"),
    ("a el administrador", "al administrador"),
    ("de el administrador", "del administrador"),
]


def insertar_despues(parrafo, texto):
    nuevo = OxmlElement("w:p")
    parrafo._p.addnext(nuevo)
    creado = Paragraph(nuevo, parrafo._parent)
    if parrafo.style is not None:
        creado.style = parrafo.style
    creado.add_run(texto)
    return creado


def main():
    doc = Document(ORIGEN)
    for parrafo in doc.paragraphs:
        for viejo, nuevo in REEMPLAZOS:
            if viejo in parrafo.text:
                for run in parrafo.runs:
                    if viejo in run.text:
                        run.text = run.text.replace(viejo, nuevo)
                if viejo in parrafo.text and parrafo.runs:
                    parrafo.runs[0].text = parrafo.text.replace(viejo, nuevo)
                    for run in parrafo.runs[1:]:
                        run.text = ""

    ancla = next(p for p in doc.paragraphs if p.text.startswith("RN-33."))
    bloques = [
        "RN-33A. Los perfiles del sistema son administrador y ayudante. Ya no existe el perfil bibliotecaria.",
        "RN-33B. Solo el administrador puede crear usuarios, asignar permisos y cambiar la contraseña, el usuario y el correo de las demás cuentas.",
        "RN-33C. El administrador reestablece su propia contraseña con tres preguntas de seguridad que él mismo registra. La contraseña de un ayudante solo la cambia el administrador.",
        "RN-33D. Al crear un ayudante se marca si puede eliminar, si puede actualizar y si puede crear. Dentro de crear se elige, por separado, libros, autores, editoriales, ubicaciones, ejemplares, lectores, préstamos e incidencias.",
        "RN-33E. El ayudante siempre ve el catálogo, los autores, las editoriales, las ubicaciones, los ejemplares, los lectores, los préstamos, las incidencias y la auditoría. Si no tiene permiso de crear, no ve el botón de alta ni el de cargar Excel. Lo mismo ocurre en ejemplares, ubicaciones, lectores y el resto de módulos: la lista se consulta, y el formulario de creación permanece oculto.",
        "RN-33F. Los botones de actualizar solo aparecen si el ayudante puede actualizar. Los de eliminar solo aparecen si puede eliminar. Cada eliminación pide confirmación.",
        "RN-33G. El módulo de usuarios solo lo ve el administrador.",
        "RN-33H. En el historial de una obra se puede seleccionar una versión y devolver la ficha a ese momento. Esa acción solo la puede hacer el administrador. Se recuperan el título, los datos bibliográficos y los autores. Los ejemplares y los préstamos no se modifican.",
    ]
    cursor = ancla
    for texto in bloques:
        cursor = insertar_despues(cursor, texto)

    for tabla in doc.tables:
        for fila in tabla.rows:
            for celda in fila.cells:
                if "bibliotecaria o ayudante" in celda.text:
                    celda.text = celda.text.replace("bibliotecaria o ayudante", "administrador o ayudante")

    doc.add_heading("6. Permisos de los perfiles", level=1)
    doc.add_paragraph(
        "Esta sección actualiza el informe inicial. El perfil bibliotecaria pasó a llamarse administrador. "
        "Los ayudantes conservan el acceso al sistema, pero cada cuenta recibe solo los permisos que el administrador marque."
    )
    doc.add_heading("6.1 Administrador", level=2)
    for texto in [
        "Ve todos los módulos, incluido Usuarios.",
        "Crea cuentas de ayudante o de administrador. Al crear otro administrador debe registrar tres preguntas y tres respuestas de seguridad.",
        "Cambia usuario, correo, contraseña y permisos de las demás cuentas.",
        "Reestablece su propia contraseña desde el ingreso, respondiendo las tres preguntas.",
        "Puede cargar un Excel con la plantilla del inventario. Las filas completas que ya estén en el catálogo se ignoran.",
        "En el historial de una obra selecciona una fila y la devuelve a esa versión.",
    ]:
        doc.add_paragraph("• " + texto)

    doc.add_heading("6.2 Ayudante", level=2)
    for texto in [
        "Siempre ve Catálogo, Autores, Editoriales, Ubicaciones, Ejemplares, Lectores, Préstamos, Incidencias y Auditoría.",
        "No ve Usuarios.",
        "Si no puede crear libros, en el catálogo no aparecen Nueva obra ni Cargar Excel con datos. Puede consultar y abrir las obras.",
        "En ejemplares, ubicaciones, lectores, autores, editoriales, préstamos e incidencias ocurre lo mismo: la lista está visible y el formulario de creación se oculta cuando no tiene ese permiso.",
        "Editar y devolver un préstamo o resolver una incidencia dependen del permiso de actualizar.",
        "Eliminar depende del permiso de eliminar y siempre pide confirmación.",
        "No puede devolver una obra a una versión anterior ni reestablecer su contraseña por su cuenta.",
    ]:
        doc.add_paragraph("• " + texto)

    doc.add_heading("6.3 Otros cambios respecto al informe inicial", level=2)
    for texto in [
        "La aplicación de uso diario es de escritorio. El acceso Biblioteca FAU.exe abre la ventana y, si hace falta, prepara Java, Node.js o la base de datos.",
        "La API escucha en el puerto 8081. La base sigue siendo PostgreSQL, con el nombre biblioteca_fau.",
        "En lectores, cédula, correo y teléfono son obligatorios. Semestre y carrera son opcionales.",
        "En usuarios se agregaron los permisos de eliminar, actualizar y crear, el detalle de qué puede crear, y las tres preguntas de seguridad.",
        "Los datos bibliográficos de una obra pueden guardarse con campos opcionales vacíos, después de confirmar cuáles faltan. Usuarios y lectores no se guardan incompletos.",
    ]:
        doc.add_paragraph("• " + texto)

    DESTINO.parent.mkdir(parents=True, exist_ok=True)
    doc.save(DESTINO)
    print(DESTINO)


if __name__ == "__main__":
    main()
